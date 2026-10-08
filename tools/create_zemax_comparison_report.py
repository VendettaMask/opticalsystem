from __future__ import annotations

import json
from collections import Counter, defaultdict
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SUMMARY_PATH = (
    ROOT
    / "artifacts"
    / "zemax-standard-samples"
    / "20261007"
    / "ra256-single-ray-fresh-import-original132"
    / "summary.json"
)
RA256_PATH = (
    ROOT
    / "artifacts"
    / "zemax-standard-samples"
    / "20261007"
    / "ra256-single-ray-fresh-import-controls6"
    / "summary.json"
)
REPORT_PATH = ROOT / "reports" / "Zemax_Workbench_Six_Lens_Comparison_Report_2026-10-08.docx"


LENS_NAMES = {
    "cooke-40-degree-field": "Cooke 40° 视场",
    "double-gauss-28-degree-field": "Double Gauss 28° 视场",
    "doublet": "Doublet",
    "even-asphere": "Even Asphere",
    "relay-lens": "Relay Lens",
    "tessar-lens-using-vignetting-factors": "Tessar 渐晕因子镜头",
}

STATUS_ORDER = ["Pass", "Close", "Difference", "Incomparable", "Error"]
STATUS_CN = {
    "Pass": "通过",
    "Close": "接近",
    "Difference": "差异",
    "Incomparable": "不可比",
    "Error": "错误",
}


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9D9D9", size="6") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_run_font(run, name="Microsoft YaHei", size=10, bold=None, color=None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def style_paragraph(paragraph, size=10, color="222222", bold=False, align=None, space_after=5, line=1.25):
    if align is not None:
        paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.line_spacing = line
    for run in paragraph.runs:
        set_run_font(run, size=size, bold=bold, color=color)


def add_body(doc: Document, text: str, bold_lead: str | None = None, keep=False):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        first = p.add_run(bold_lead)
        first.bold = True
        p.add_run(text[len(bold_lead) :])
    else:
        p.add_run(text)
    p.paragraph_format.keep_together = keep
    style_paragraph(p, size=10.5, space_after=6, line=1.35)
    return p


def add_bullet(doc: Document, text: str, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.add_run(text)
    style_paragraph(p, size=10.2, space_after=3, line=1.25)
    return p


def add_heading(doc: Document, text: str, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(5)
    size = 15 if level == 1 else 12
    for run in p.runs:
        set_run_font(run, size=size, bold=True, color="000000")
    return p


def add_table(doc: Document, headers, rows, widths=None, font_size=8.6, center_cols=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    header = table.rows[0]
    set_repeat_table_header(header)
    for idx, text in enumerate(headers):
        cell = header.cells[idx]
        cell.text = str(text)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, "1F4E78")
        set_cell_margins(cell, top=100, bottom=100)
        if widths:
            cell.width = widths[idx]
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            for run in p.runs:
                set_run_font(run, size=font_size, bold=True, color="FFFFFF")
    for r_idx, row_values in enumerate(rows):
        row = table.add_row()
        if r_idx % 2:
            for cell in row.cells:
                set_cell_shading(cell, "F2F6FA")
        for idx, value in enumerate(row_values):
            cell = row.cells[idx]
            cell.text = "" if value is None else str(value)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if widths:
                cell.width = widths[idx]
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center_cols and idx in center_cols else WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.1
                for run in p.runs:
                    set_run_font(run, size=font_size, color="222222")
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    return table


def fmt_num(value):
    if value is None:
        return "—"
    if value == 0:
        return "0"
    av = abs(value)
    if av < 0.001 or av >= 1000:
        return f"{value:.3E}"
    return f"{value:.6f}".rstrip("0").rstrip(".")


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.3

    title = styles["Title"]
    title.font.name = "Microsoft YaHei"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title_ppr = title._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for style_name in ("Heading 1", "Heading 2", "Heading 3"):
        style = styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.color.rgb = RGBColor(0, 0, 0)

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Zemax 与 Workbench 六镜头结果对比报告  |  2026-10-08")
    style_paragraph(p, size=8, color="666666", space_after=0, line=1.0)


def build_report() -> Path:
    summary = json.loads(SUMMARY_PATH.read_text(encoding="utf-8"))
    ra256 = json.loads(RA256_PATH.read_text(encoding="utf-8"))
    rows = summary["rows"]

    by_lens = defaultdict(Counter)
    non_pass_by_lens = defaultdict(list)
    by_analysis = defaultdict(Counter)
    for item in rows:
        lens = item["lens"]
        conclusion = item["currentConclusion"]
        by_lens[lens][conclusion] += 1
        if conclusion != "Pass":
            non_pass_by_lens[lens].append(item)
            by_analysis[item["key"]][conclusion] += 1

    doc = Document()
    configure_document(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(34)
    title.paragraph_format.space_after = Pt(10)
    title.add_run("Zemax 与 Workbench 六镜头结果对比报告")
    title_ppr = title._p.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("OpticStudio 2026 R1 冻结基准与正式共享 Core 复算")
    style_paragraph(subtitle, size=13, color="333333", space_after=8, line=1.2)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run("报告日期 2026年10月8日   数据截止 2026年10月7日")
    style_paragraph(meta, size=9.5, color="666666", space_after=22, line=1.1)

    add_body(
        doc,
        "本报告比较六份官方顺序模式镜头在 Zemax OpticStudio 2026 R1 与 Optical System Design Workbench 正式共享 Core 中的分析结果。当前 132 项原设置比较为 95 项通过、12 项接近、16 项差异、8 项不可比和 1 项错误；独立 RMS 控制 23/23 通过，但总体发布门禁尚未通过。",
        keep=True,
    )

    add_table(
        doc,
        ["指标", "结果", "解释"],
        [
            ["原设置比较", "95 / 132 通过", "通过率 72.0%，两条复算路径结论一致"],
            ["独立 RMS 控制", "23 / 23 通过", "此前 17 项与新增 RA256 6 项分开计数"],
            ["首批官方镜头", "6 / 27", "其余 21 份尚未运行"],
            ["发布判断", "未通过", "仍有衍射、Huygens、MTF、照度及不可比项目"],
        ],
        widths=[Cm(4.1), Cm(3.7), Cm(8.6)],
        font_size=9.1,
        center_cols={1},
    )

    doc.add_page_break()

    add_heading(doc, "1 对比范围与判定方法", 1)
    add_body(
        doc,
        "对比对象为 Cooke 40° 视场、Double Gauss 28° 视场、Doublet、Even Asphere、Relay Lens 和带渐晕因子的 Tessar。每个镜头包含 22 项原分析设置，共 132 项。原生版本、许可、请求、分析设置、单位、采样和容差均随冻结证据保存。",
    )
    add_body(
        doc,
        "两条复算路径分别使用不改写的历史快照和从原始 ZMX 重新导入的当前模型。两条路径使用相同的原生数组、请求和容差；报告不把两条路径重复计数，也不把 23 项独立 RMS 控制计入 132 项矩阵。",
    )
    add_table(
        doc,
        ["分类", "含义"],
        [
            ["Pass 通过", "误差满足该分析的冻结容差和覆盖要求。"],
            ["Close 接近", "超过通过门槛，但仍处于单独记录的接近范围。"],
            ["Difference 差异", "存在可量化且超过容差的数值差异。"],
            ["Incomparable 不可比", "原生或 Workbench 输出缺少可按同一物理量比较的有效数组。"],
            ["Error 错误", "当前请求无法形成有效比较；不能用填零或伪造数组替代。"],
        ],
        widths=[Cm(4.2), Cm(12.2)],
        font_size=9.1,
    )
    add_body(
        doc,
        "证据边界：GB 级完整 RA256 逐光线数组没有全部提交到仓库，仓库保留其清单、SHA-256、统计摘要和验证结果。已提交的普通捕获、RMS 曲线、单光线报告及较小逐光线审计足以复核本报告所列结论。",
        bold_lead="证据边界：",
    )

    add_heading(doc, "2 总体结果与演进", 1)
    progression = [
        ["首批比较", "76", "15", "30", "10", "1"],
        ["OPD 修复", "82", "12", "27", "10", "1"],
        ["渐晕与 RMS 采样修复", "84", "12", "27", "8", "1"],
        ["Tessar RMS 修复", "85", "12", "26", "8", "1"],
        ["RA256 与单光线修复", "95", "12", "16", "8", "1"],
    ]
    add_table(
        doc,
        ["阶段", "Pass", "Close", "Difference", "Incomparable", "Error"],
        progression,
        widths=[Cm(5.3), Cm(2.0), Cm(2.0), Cm(2.5), Cm(2.9), Cm(1.7)],
        font_size=8.7,
        center_cols={1, 2, 3, 4, 5},
    )
    add_body(
        doc,
        "从首批比较到当前阶段，Pass 从 76 项增至 95 项，Difference 从 30 项降至 16 项。最终阶段相对上一阶段新增 10 项 Pass，且既有 85 项 Pass 无回退。改进主要来自自动 STOP 导入语义、无限物距单光线首段路径、OPD 约定、渐晕接纳和 RMS 采样链路修正。",
    )

    lens_rows = []
    for lens in LENS_NAMES:
        counter = by_lens[lens]
        pass_rate = counter["Pass"] / 22 * 100
        lens_rows.append(
            [
                LENS_NAMES[lens],
                counter["Pass"],
                counter["Close"],
                counter["Difference"],
                counter["Incomparable"],
                counter["Error"],
                f"{pass_rate:.1f}%",
            ]
        )
    add_table(
        doc,
        ["镜头", "Pass", "Close", "Difference", "Incomp.", "Error", "通过率"],
        lens_rows,
        widths=[Cm(5.1), Cm(1.7), Cm(1.7), Cm(2.2), Cm(2.0), Cm(1.5), Cm(2.2)],
        font_size=8.5,
        center_cols={1, 2, 3, 4, 5, 6},
    )

    add_heading(doc, "3 逐镜头结论", 1)
    lens_notes = {
        "cooke-40-degree-field": "主要差异集中在衍射包围能量、Huygens Through Focus MTF、相对照度、PSF 和 Huygens PSF 截面；Wavefront 为 Close。单光线路径与 OPD 已转为 Pass。",
        "double-gauss-28-degree-field": "主要差异集中在 Huygens Through Focus MTF、相对照度、PSF 和 Huygens PSF 截面；衍射包围能量不可比。单光线路径与 OPD 已转为 Pass。",
        "doublet": "总体 16/22 通过。衍射包围能量仍有差异，RMS vs Field 与相对照度不可比；Field Curvature and Distortion 因单视场参考映射奇异保留为 Error。",
        "even-asphere": "主要差异在 Huygens Through Focus MTF、Huygens PSF、Huygens PSF 截面和 MTF；衍射包围能量与 PSF 不可比。",
        "relay-lens": "六镜头中通过率最高，为 18/22。剩余 Difference 为衍射包围能量和 Huygens Through Focus MTF；PSF 不可比，Huygens PSF 截面为 Close。",
        "tessar-lens-using-vignetting-factors": "当前无 Difference 或 Error，17/22 通过；衍射包围能量和 PSF 不可比，相对照度、Huygens PSF 截面与 Wavefront 为 Close。自动 STOP 修复后 RA256 接纳和首次截断差异归零。",
    }
    for lens in LENS_NAMES:
        add_heading(doc, LENS_NAMES[lens], 2)
        add_body(doc, lens_notes[lens])

    add_heading(doc, "4 已确认修复与独立控制", 1)
    add_table(
        doc,
        ["修复或控制", "结果", "结论"],
        [
            ["自动 STOP 导入", "46 条接纳差异与 4 条首次截断差异归零", "自动 DIAM 保留估值语义，不再误建实体圆孔"],
            ["Single Ray Trace 首段", "10 项原 Difference 转为 Pass", "无限物距首段改用第一面局部参考平面"],
            ["独立 RMS 控制", "23 / 23 Pass", "17 项历史控制与 6 项 RA256 控制均通过"],
            ["完整边缘光瞳", "308,808 条输入与 2,521,932 个逐面结果", "接纳、截断、传播和坐标分别核验"],
            ["大型 JSON", "1.37 GB Core JSON 流式写入完成", "消除整段 UTF-16 字符串分配导致的 OOM"],
        ],
        widths=[Cm(4.3), Cm(5.0), Cm(7.1)],
        font_size=8.7,
    )

    ra_rows = []
    for item in ra256["rows"]:
        lens_key = item["lens"].replace("-ra-256-remove", "").replace("-ra-256-retain", "")
        mode = "保留渐晕因子" if item["lens"].endswith("retain") else "移除渐晕因子"
        ra_rows.append(
            [
                LENS_NAMES.get(lens_key, lens_key),
                mode,
                STATUS_CN[item["currentConclusion"]],
                fmt_num(item["currentWorstMaxAbsolute"]),
            ]
        )
    add_heading(doc, "RA256 独立 RMS 控制", 2)
    add_table(
        doc,
        ["镜头", "模式", "结果", "最大绝对误差 μm"],
        ra_rows,
        widths=[Cm(5.0), Cm(4.0), Cm(2.2), Cm(5.2)],
        font_size=8.8,
        center_cols={2, 3},
    )
    add_body(
        doc,
        "六项 RA256 控制均为 Pass。需要注意，Tessar 的原生 RMS 从 RA128 到 RA256 仍变化约 0.164328 μm，且变化并非单调，因此相同网格下的数值一致不能替代更高密度收敛证明。Relay 的有限共轭瞄准坐标残差约为 1.115×10⁻⁷ mm，仍是后续精度工作。",
    )

    add_heading(doc, "5 未完成问题与优先级", 1)
    analysis_rows = []
    for key, counter in sorted(by_analysis.items(), key=lambda kv: (-sum(kv[1].values()), kv[0])):
        analysis_rows.append(
            [
                key,
                counter["Difference"],
                counter["Close"],
                counter["Incomparable"],
                counter["Error"],
                sum(counter.values()),
            ]
        )
    add_table(
        doc,
        ["分析类型", "Difference", "Close", "Incomp.", "Error", "未通过合计"],
        analysis_rows,
        widths=[Cm(6.2), Cm(2.2), Cm(1.8), Cm(2.0), Cm(1.5), Cm(2.7)],
        font_size=8.3,
        center_cols={1, 2, 3, 4, 5},
    )

    add_body(doc, "建议按以下顺序推进：", bold_lead="建议按以下顺序推进：")
    add_bullet(doc, "P0：完成 Tessar RA512 或更高密度收敛研究，并收紧 Relay 有限共轭瞄准精度。")
    add_bullet(doc, "P0：优先修正 Huygens Through Focus MTF、Huygens PSF 截面、衍射包围能量和 PSF。")
    add_bullet(doc, "P1：处理相对照度、普通 MTF、Huygens PSF 与 Wavefront 的剩余差异或 Close。")
    add_bullet(doc, "P1：为 Doublet 单视场畸变奇异映射建立明确的不可比或有效参考规则，不填造数组。")
    add_bullet(doc, "P2：按相同捕获纪律扩展剩余 21 份官方镜头，再执行完整发布门禁。")

    add_heading(doc, "6 结论", 1)
    add_body(
        doc,
        "当前结果证明共享 Core 在几何光线、OPD、RMS 场曲线和光阑接纳方面已形成可靠的高精度基线，尤其是 23 项独立 RMS 控制全部通过，自动 STOP 与单光线路径问题得到可重复修复。",
    )
    add_body(
        doc,
        "但 132 项原设置中仍有 37 项不是 Pass，且只覆盖计划中 27 份官方镜头的 6 份。主要缺口集中在衍射和 Huygens 系列、照度、部分 MTF，以及不可比较的输出契约。因此当前状态适合继续定向数值收口，不适合宣称已完成全面 Zemax 等价或商业发布认证。",
    )

    appendix_a_heading = add_heading(doc, "附录 A 非 Pass 项明细", 1)
    appendix_a_heading.paragraph_format.space_after = Pt(12)
    detail_rows = []
    for lens in LENS_NAMES:
        for item in non_pass_by_lens[lens]:
            detail_rows.append(
                [
                    LENS_NAMES[lens],
                    item["key"],
                    STATUS_CN[item["currentConclusion"]],
                    fmt_num(item["currentWorstNrmse"]),
                    fmt_num(item["currentWorstMaxAbsolute"]),
                ]
            )
    add_table(
        doc,
        ["镜头", "分析类型", "分类", "最差 NRMSE", "最大绝对误差"],
        detail_rows,
        widths=[Cm(4.3), Cm(5.3), Cm(2.0), Cm(2.3), Cm(2.5)],
        font_size=7.8,
        center_cols={2, 3, 4},
    )

    appendix_b_heading = add_heading(doc, "附录 B 证据索引", 1)
    appendix_b_heading.paragraph_format.space_after = Pt(10)
    evidence = [
        "docs/validation/ZEMAX_RA256_SINGLE_RAY_REPAIR_2026-10-07.json",
        "docs/ZEMAX_RA256_SINGLE_RAY_REPAIR_2026-10-07.md",
        "artifacts/zemax-standard-samples/20261007/ra256-single-ray-fresh-import-original132/summary.json",
        "artifacts/zemax-standard-samples/20261007/ra256-single-ray-old-snapshot-original132/summary.json",
        "validation/zemax/2026-r1/ra256-field-sampling-2026-10-07/manifest.json",
        "validation/zemax/2026-r1/single-ray-path-2026-10-07/manifest.json",
        "docs/validation/ZEMAX_STANDARD_SAMPLE_PLAN_2026-10-06.json",
    ]
    for path in evidence:
        p = doc.add_paragraph()
        p.add_run(path)
        style_paragraph(p, size=8.5, color="444444", space_after=3, line=1.15)

    REPORT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(REPORT_PATH)
    return REPORT_PATH


if __name__ == "__main__":
    print(build_report())
